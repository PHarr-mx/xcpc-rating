"""staged 导入 core 层流程：stage → confirm / discard（盲区补齐，不依赖 Reflex）。"""

import shutil
from datetime import date
from pathlib import Path

import pytest
from sqlalchemy import select

from xcpc_core.db.tables import Contest, ImportBatch, Player, Team
from xcpc_core.importer.api import (
    confirm_import_batch,
    discard_import_batch,
    stage_formal_xlsx,
)
from xcpc_core.player.store import find_repo_root

_HEADERS = ["Rank", "Organization", "Team", "Member1", "Member2", "Member3",
            "Solved", "Penalty", "Organization Rank", "Medal",
            "A", "B", "C", "D", "E", "F", "G", "H"]  # 题号列供 total_problems 统计


def _write_xlsx(path: Path) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws_all = wb.active
    ws_all.title = "所有队伍"
    ws_all["A1"] = "核心测试赛"
    ws_all.append(_HEADERS)
    ws_all.append([1, "西南民族大学", "一队", "张三", "李四", None, 5, 100, 1, "gold"])
    ws_all.append([2, "西南民族大学", "二队", "王五", None, None, 3, 200, 2, "silver"])
    ws_all.append([3, "外校大学", "客队", "外人", None, None, 9, 50, None, "gold"])
    ws_formal = wb.create_sheet("正式组")
    ws_formal["A1"] = "核心测试赛"
    ws_formal.append(_HEADERS)
    ws_formal.append([1, "西南民族大学", "一队", "张三", "李四", None, 5, 100, 1, "gold"])
    ws_formal.append([2, "西南民族大学", "二队", "王五", None, None, 3, 200, 2, "silver"])
    wb.save(path)


@pytest.fixture
def repo(tmp_path):
    """临时仓库：需要 contest_weights.yaml（权重查表）与 raw 目录（归档）。"""
    real_root = find_repo_root()
    (tmp_path / "data/config").mkdir(parents=True)
    (tmp_path / "data/raw/formal").mkdir(parents=True)
    shutil.copy(real_root / "data/config/contest_weights.yaml", tmp_path / "data/config/contest_weights.yaml")
    shutil.copy(real_root / "data/config/school.yaml", tmp_path / "data/config/school.yaml")
    return tmp_path


@pytest.fixture
def xlsx(repo):
    path = repo / "uploads_core_test.xlsx"
    _write_xlsx(path)
    return path


def _params(contest_id: str):
    from xcpc_core.importer.models import FormalImportParams

    return FormalImportParams(
        contest_id=contest_id,
        date=date(2026, 5, 18),
        contest_type="icpc_provincial",
        school_organizations=["西南民族大学"],
    )


def test_stage_then_confirm_writes_everything(db_session, repo, xlsx):
    summary = stage_formal_xlsx(
        xlsx, _params("core_staged_c1"),
        uploaded_by=1, filename=xlsx.name, repo_root=repo, session=db_session,
    )
    assert summary.status == "staged"
    assert summary.total_teams == 3  # 所有队伍页 3 行（含外校）
    assert summary.standings_count == 2  # 本校 2 队
    assert {u.name for u in summary.unmatched_players} == {"张三", "李四", "王五"}
    assert db_session.scalars(select(Player)).all() == []  # 解析阶段不写正式表

    result = confirm_import_batch(
        summary.batch_id,
        {"张三": "new", "李四": "new", "王五": "new"},
        repo_root=repo, session=db_session,
    )
    assert result.standings_imported == 2
    assert {p.name for p in result.players_created} == {"张三", "李四", "王五"}

    players = {p.id: p.name for p in db_session.scalars(select(Player)).all()}
    assert players == {"p001": "张三", "p002": "李四", "p003": "王五"}
    assert db_session.scalars(select(Team)).all() != []
    contest = db_session.get(Contest, "core_staged_c1")
    assert contest is not None and contest.source_type == "formal"
    (batch,) = db_session.scalars(select(ImportBatch)).all()
    assert batch.status == "confirmed"
    # raw 归档写入临时仓库（而非真实仓库）
    assert (repo / "data/raw/formal/core_staged_c1.json").is_file()


def test_confirm_twice_rejected(db_session, repo, xlsx):
    summary = stage_formal_xlsx(
        xlsx, _params("core_staged_c2"),
        uploaded_by=1, repo_root=repo, session=db_session,
    )
    confirm_import_batch(summary.batch_id, {"张三": "new", "李四": "new", "王五": "new"},
                         repo_root=repo, session=db_session)
    with pytest.raises(ValueError, match="不能确认"):
        confirm_import_batch(summary.batch_id, {}, repo_root=repo, session=db_session)


def test_discard_then_confirm_rejected(db_session, repo, xlsx):
    summary = stage_formal_xlsx(
        xlsx, _params("core_staged_c3"),
        uploaded_by=1, repo_root=repo, session=db_session,
    )
    discard_import_batch(summary.batch_id, session=db_session)
    (batch,) = db_session.scalars(select(ImportBatch)).all()
    assert batch.status == "discarded"
    with pytest.raises(ValueError, match="不能确认"):
        confirm_import_batch(summary.batch_id, {"张三": "new"}, repo_root=repo, session=db_session)
